import os
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

# =============================================================
# 1. 공통 부모 클래스 (BaseReasonHandler)
# =============================================================
class BaseReasonHandler:
    def __init__(self, items, log_key, mapped_reason=None):
        self.items = items
        self.log_key = log_key
        self.full_text = " ".join(items)
        self.full_text_no_space = self.full_text.replace(" ", "")
        
        self.reason = mapped_reason if mapped_reason else (items[2].strip() if len(items) > 2 else "")

    def extract_basics(self):
        cpr = "기타(O)"
        if len(self.items) > 0:
            clean = self.items[0].replace(" ", "")
            if clean == '고객': cpr = '고객'
            elif clean == '파트너': cpr = '파트너'
            elif clean == '라이더': cpr = '라이더'

        service = "기타(O)"
        if len(self.items) > 1:
            clean = self.items[1].replace(" ", "")
            if clean == 'B마트': service = 'B마트'
            elif clean == '장보기': service = '장보기'
            elif clean == '푸드': service = '푸드'

        request = self.items[3].strip() if len(self.items) > 3 else ""
        if len(self.items) > 5:
            process = " / ".join(self.items[4:-1]).strip()
            result_raw = self.items[-1].strip()
        elif len(self.items) == 5:
            process = ""
            result_raw = self.items[4].strip()
        else:
            process = ""
            result_raw = ""

        # 여부 플래그
        three_way = "O" if "파트너 " in process or "고객" in process else "X"
        misroute = "O" if "고객오인입" in self.full_text_no_space else "X"
        
        # ★ 문제상황여부 플래그
        problem_situation = "O" if any(kw in self.full_text_no_space for kw in ["불만", "미수긍", "항의"]) else "X"

        clean_consent = re.sub(r'제공\s*동의', '', self.full_text)
        manual_class = "O"
        result = result_raw
        if "동의" in clean_consent or "수긍" in clean_consent:
            result = request
            manual_class = "X"

        advance = "O" if "선조치" in self.full_text_no_space else "X"
        coupon = "O" if "쿠폰" in self.full_text_no_space else "X"
        transfer = "O" if "이관" in self.full_text_no_space else "X"

        return cpr, service, request, process, result, three_way, misroute, problem_situation, manual_class, advance, coupon, transfer

    def get_subclass(self, cpr, service, request, process):
        return "예외처리"

    def process_log(self):
        cpr, service, request, process, result, \
        three_way, misroute, problem_situation, manual_class, advance, coupon, transfer = self.extract_basics()
        
        processed_result = self.get_subclass(cpr, service, request, process)

        return {
            "순번": self.log_key,
            "대상(CPR)": cpr,
            "서비스(BSFO)": service,
            "처리결과": processed_result,       
            "인입이유": self.reason,
            "요청사항": request,
            "처리과정": process,
            "원문_처리결과": result,           
            "3자통화여부": three_way,
            "고객오인입여부": misroute,
            "문제상황여부": problem_situation,
            "수동분류필요": manual_class,
            "선조치여부": advance,
            "쿠폰발급여부": coupon,
            "이관여부": transfer,
            "원본내용": " / ".join(self.items)
        }

# =============================================================
# 2. 인입 이유별 자식 클래스 (비즈니스 룰 기반 세부 로직)
# =============================================================
class StockOutHandler(BaseReasonHandler):
    def get_subclass(self, cpr, service, request, process):
        if "고객부재" in self.full_text_no_space:
            if "수기" in self.full_text: return "수기 환불"
            elif "부분취소" in self.full_text_no_space or "현금" in self.full_text or "계좌" in self.full_text or "부분" in self.full_text: return "부분 취소"
            elif "주문" in self.full_text: return "주문 취소"
            return "예외처리"
        else:
            if "수기" in self.full_text: return "수기 환불"
            elif "대체동의" in self.full_text_no_space or "대체상품동의" in self.full_text_no_space: return "대체"
            elif "부분취소" in self.full_text_no_space or "현금" in self.full_text or "계좌" in self.full_text or "환불" in self.full_text: return "부분 취소"
            elif "주문" in self.full_text: return "주문 취소"
            return "예외처리"

class MissingHandler(BaseReasonHandler):
    def get_subclass(self, cpr, service, request, process):
        text_ns = self.full_text_no_space
        
        if any(kw in text_ns for kw in ["분리배달", "상담중", "찾음", "정상"]): return "누락 X"
        if any(kw in text_ns for kw in ["앤써", "앱접수", "슬랙", "계좌", "수기"]): return "환불 O"
        if any(kw in text_ns for kw in ["추가배차", "재배차", "추가배달비"]): return "교환/재배달 O"
            
        if "부재" in self.full_text: return "부재"
        clean_consent = re.sub(r'제공\s*동의', '', self.full_text)
        is_consent = "O" if "동의" in clean_consent else "X"
        
        if "환불" in self.full_text: return f"환불 {is_consent}"
        elif "교환" in self.full_text or "재배달" in self.full_text: return f"교환/재배달 {is_consent}"
        return "예외처리"

class MisdeliveryHandler(BaseReasonHandler):
    def get_subclass(self, cpr, service, request, process):
        text_ns = self.full_text_no_space
        if "직접픽업불가" in text_ns:
            if "라이더정정배달" in text_ns: return "직접 픽업 불가 - 라이더 정정배달 O"
            return "직접 픽업 불가 - 주문 취소"
        elif "직접픽업" in text_ns: return "직접 픽업"
        elif "정정배달" in text_ns or "상담중" in text_ns: return "정정 배달"
        elif "자체배달" in text_ns: return "자체 배달"
        elif "주문취소" in text_ns: return "주문 취소"
        return "예외처리"

class DeliveryDelayHandler(BaseReasonHandler):
    def get_subclass(self, cpr, service, request, process):
        if "수긍" in self.full_text: return "수긍 O"
        if "주문취소" in self.full_text_no_space: return "수긍 X - 주문 취소 O"
        return "수긍 X - 주문 취소 X"

class DispatchDelayHandler(BaseReasonHandler):
    def get_subclass(self, cpr, service, request, process):
        text_ns = self.full_text_no_space
        if "상담중" in text_ns: return "해결"
        if "취소" in text_ns: return "주문 취소"
        if "긴급배차" in text_ns: return "긴급 배차"
        if "위치안내" in text_ns: return "라이더 위치 안내"
        return "예외처리"

class QualityDefectHandler(BaseReasonHandler):
    def get_subclass(self, cpr, service, request, process):
        if any(kw in self.full_text_no_space for kw in ["앱접수", "수기", "반품"]): return "환불 O"
            
        clean_consent = re.sub(r'제공\s*동의', '', self.full_text)
        is_consent = "O" if "동의" in clean_consent else "X"

        if service == "B마트":
            if "환불" in self.full_text: return "환불 O"
            if "재배달" in self.full_text or "교환" in self.full_text: return "재배달/교환 O"
            return "예외처리"
        elif service == "장보기":
            if "환불" in self.full_text: return f"환불 {is_consent}"
            if "재배달" in self.full_text or "교환" in self.full_text: return f"재배달/교환 {is_consent}"
            return "예외처리"
            
        return "예외처리"

class ChangeOfMindHandler(BaseReasonHandler):
    def get_subclass(self, cpr, service, request, process):
        if "주문취소" in request.replace(" ", ""):
            clean_consent = re.sub(r'제공\s*동의', '', self.full_text)
            if "동의" in clean_consent: return "주문 취소 O"
            if "파트너부재" in self.full_text_no_space: return "주문 취소 - 파트너 부재"
            return "예외처리"
        return "예외처리"

# =============================================================
# 3. 라우터
# =============================================================
def route_and_process(items, log_key):
    reason_raw = items[2] if len(items) > 2 else ""
    reason_clean = reason_raw.replace(" ", "")
    
    if any(k in reason_clean for k in ["불량", "저하", "곰팡이", "이물질", "선도", "변질"]):
        handler = QualityDefectHandler(items, log_key, mapped_reason="품질 불량")
    elif "재고소진" in reason_clean: handler = StockOutHandler(items, log_key, mapped_reason="재고 소진")
    elif "누락" in reason_clean: handler = MissingHandler(items, log_key, mapped_reason="누락")
    elif "오배달" in reason_clean: handler = MisdeliveryHandler(items, log_key, mapped_reason="오배달")
    elif "배달지연" in reason_clean: handler = DeliveryDelayHandler(items, log_key, mapped_reason="배달 지연")
    elif "배차지연" in reason_clean: handler = DispatchDelayHandler(items, log_key, mapped_reason="배차 지연")
    elif "오포장" in reason_clean: handler = QualityDefectHandler(items, log_key, mapped_reason="오포장")
    elif "단순변심" in reason_clean: handler = ChangeOfMindHandler(items, log_key, mapped_reason="단순 변심")
    elif "파손" in reason_clean: handler = QualityDefectHandler(items, log_key, mapped_reason="파손")
    else: 
        handler = BaseReasonHandler(items, log_key)

    return handler.process_log()

# =============================================================
# 4. 메인 
# =============================================================
if __name__ == "__main__":
    current_dir = os.path.dirname(os.path.abspath(__file__))
    input_file_path = os.path.join(current_dir, 'data', 'consult_log.txt')
    if not os.path.exists(input_file_path):
        alt_path = os.path.join(current_dir, 'data', 'consult_log')
        if os.path.exists(alt_path): input_file_path = alt_path

    output_dir = os.path.join(current_dir, 'outputs')
    os.makedirs(output_dir, exist_ok=True)
    output_excel_path = os.path.join(output_dir, 'reason_class_data.xlsx')

    detailed_records = []
    valid_count = 0
    with open(input_file_path, 'r', encoding='utf-8-sig') as f:
        for raw_line in f:
            line = raw_line.strip()
            
            if not line or line.startswith('2026.'): continue
            cleaned_line = re.sub(r'^[\[\(]?\d+[\.\)\-\]\s]+', '', line).strip()
            if not re.search(r'[가-힣a-zA-Z]', cleaned_line): continue
            
            normalized_line = cleaned_line.replace('／', '/')
            items = [item.strip() for item in normalized_line.split('/')]
            
            valid_count += 1
            log_key = f"log_{valid_count}"
            
            detailed_records.append(route_and_process(items, log_key))

    df = pd.DataFrame(detailed_records)
    TARGET_REASONS = ["품질 불량", "재고 소진", "누락", "오배달", "배달 지연", "배차 지연", "오포장", "단순 변심", "파손"]
    
    # ★ 3. 통계 요약 시트 고도화 (처리 결과별로 문제상황/쿠폰 건수 및 비율 산출)
    summary_records = []
    all_sheet_names = TARGET_REASONS + ["작업 필요"]
    
    for sheet_name in all_sheet_names:
        if sheet_name == "작업 필요":
            sub_df = df[~df['인입이유'].isin(TARGET_REASONS)]
        else:
            sub_df = df[df['인입이유'] == sheet_name]
            
        total_logs_in_sheet = len(sub_df)
        if total_logs_in_sheet == 0:
            continue
            
        val_counts = sub_df['처리결과'].value_counts()
        for res, count in val_counts.items():
            # 특정 처리결과의 데이터만 필터링
            res_df = sub_df[sub_df['처리결과'] == res]
            
            # 특정 처리결과 안에서의 문제상황 / 쿠폰발급 건수
            problem_count = len(res_df[res_df['문제상황여부'] == 'O'])
            coupon_count = len(res_df[res_df['쿠폰발급여부'] == 'O'])
            
            # 비율 계산
            res_ratio = (count / total_logs_in_sheet) * 100
            prob_ratio = (problem_count / count) * 100 if count > 0 else 0
            coup_ratio = (coupon_count / count) * 100 if count > 0 else 0

            summary_records.append({
                "인입이유(시트명)": sheet_name,
                "시트 내 총 건수": total_logs_in_sheet,
                "처리결과": res,
                "해당 결과 건수": count,
                "처리결과 비율(%)": f"{res_ratio:.1f}%",
                "문제상황 건수": problem_count,
                "문제상황 비율(%)": f"{prob_ratio:.1f}%",
                "쿠폰발급 건수": coupon_count,
                "쿠폰발급 비율(%)": f"{coup_ratio:.1f}%"
            })
            
    summary_df = pd.DataFrame(summary_records)

    grouping_cols = [
        "대상(CPR)", "서비스(BSFO)", "처리결과", "인입이유", 
        "요청사항", "처리과정", "원문_처리결과", "3자통화여부", 
        "고객오인입여부", "문제상황여부", "수동분류필요", "선조치여부", "쿠폰발급여부", "이관여부", "원본내용"
    ]
    
    final_cols_order = ["순번", "대상(CPR)", "서비스(BSFO)", "처리결과"] + \
                       [c for c in grouping_cols if c not in ["대상(CPR)", "서비스(BSFO)", "처리결과"]]

    with pd.ExcelWriter(output_excel_path, engine='openpyxl') as writer:
        
        summary_df.to_excel(writer, sheet_name="통계 요약", index=False)
        
        sheet_dict = {reason: df[df['인입이유'] == reason] for reason in TARGET_REASONS}
        sheet_dict["작업 필요"] = df[~df['인입이유'].isin(TARGET_REASONS)]

        for sheet_name, sheet_df in sheet_dict.items():
            if sheet_df.empty:
                continue
                
            grouped_df = sheet_df.groupby(grouping_cols, dropna=False).agg({
                '순번': lambda x: ', '.join(x)
            }).reset_index()
            
            grouped_df['log_count'] = grouped_df['순번'].apply(lambda x: len(str(x).split(',')))
            grouped_df['서비스_sort'] = pd.Categorical(
                grouped_df['서비스(BSFO)'], 
                categories=['B마트', '장보기', '푸드', '기타(O)'], ordered=True)
            grouped_df['대상_sort'] = pd.Categorical(
                grouped_df['대상(CPR)'], 
                categories=['고객', '파트너', '라이더', '기타(O)'], ordered=True)

            grouped_df = grouped_df[final_cols_order + ['log_count', '서비스_sort', '대상_sort']]
            grouped_df.sort_values(
                by=['log_count', '서비스_sort', '대상_sort', '처리과정'], 
                ascending=[False, True, True, True], 
                inplace=True
            )
            
            grouped_df.drop(columns=['log_count', '서비스_sort', '대상_sort'], inplace=True)
            grouped_df.to_excel(writer, sheet_name=sheet_name, index=False)

    wb = load_workbook(output_excel_path)
    
    dv_cpr = DataValidation(type="list", formula1='"고객,파트너,라이더,기타(O)"', allow_blank=True)
    dv_bsfo = DataValidation(type="list", formula1='"B마트,장보기,푸드,기타(O)"', allow_blank=True)
    dv_ox = DataValidation(type="list", formula1='"O,X"', allow_blank=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        if sheet_name == "통계 요약":
            ws.auto_filter.ref = ws.dimensions
            continue
            
        max_row = ws.max_row
        ws.add_data_validation(dv_cpr)
        ws.add_data_validation(dv_bsfo)
        ws.add_data_validation(dv_ox)

        dv_cpr.add(f"B2:B{max_row}")
        dv_bsfo.add(f"C2:C{max_row}")
        # I열부터 O열까지 여부 플래그 (3자통화, 오인입, 문제상황, 수동분류, 선조치, 쿠폰, 이관)
        dv_ox.add(f"I2:O{max_row}")

        ws.auto_filter.ref = ws.dimensions

    wb.save(output_excel_path)

    print("\n" + "=" * 65)
    print("   📊 요약 시트 상세화 (문제상황/쿠폰 건수 및 비율 명시) 적용 완료")
    print("=" * 65)
    print("✔ '통계 요약' 시트에 각 처리결과별 문제 상황 / 쿠폰 발생 건수 추가")
    print("✔ 각 결과값 안에서 문제상황이 차지하는 비중(%) 정밀 계산 완료")
    print("=" * 65 + "\n")